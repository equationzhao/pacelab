"""Extract only measured exercise lactate-power pairs; never interpolate readings."""
import hashlib
import json
from pathlib import Path
import zipfile

import openpyxl

HERE = Path(__file__).resolve().parent
manifest = json.loads((HERE / "sources.json").read_text())
for item in manifest["files"]:
    data = (HERE / "sources" / item["name"]).read_bytes()
    assert hashlib.sha256(data).hexdigest() == item["sha256"], item["name"]

curves = []
with zipfile.ZipFile(HERE / "sources/nolte-raw.zip") as archive:
    for name in sorted(archive.namelist()):
        if not name.endswith(".txt"):
            continue
        rows = []
        for row, line in enumerate(archive.read(name).decode().splitlines(), 1):
            if not line.strip():
                continue
            power, lactate = map(float, line.split())
            rows.append(dict(power=power, lactate=lactate, quality="normal", sourceRow=row))
        curves.append(dict(id="N2022-" + Path(name).stem, source="Nolte2022",
                           subject=Path(name).stem, file=name,
                           rest=[p for p in rows if p["power"] == 0],
                           points=[p for p in rows if p["power"] > 0]))

workbook = openpyxl.load_workbook(HERE / "sources/pallares-2016.xlsx", read_only=True, data_only=True)
for sheet in workbook:
    if "IMA" not in sheet.title:
        continue
    # Seven sheets have an additional duplicate power column B.
    assert sheet.max_column in (123, 124)
    lactate_column = 4 if sheet.max_column == 124 else 3
    rows = []
    for row_number, row in enumerate(sheet.values, 1):
        power, lactate = row[0], row[lactate_column - 1]
        if not isinstance(power, (int, float)) or not isinstance(lactate, (int, float)):
            continue
        if sheet.max_column == 124:
            assert row[1] == power
        rows.append(dict(power=power, lactate=lactate, quality="normal",
                         sourceRow=row_number, lactateColumn=lactate_column))
    curves.append(dict(id="P2016-" + sheet.title.replace(" ", "-"), source="Pallares2016",
                       subject=sheet.title.split()[0], sheet=sheet.title,
                       rest=[p for p in rows if p["power"] == 0],
                       points=[p for p in rows if p["power"] > 0]))
assert len(curves) == 52
for c in curves:
    assert all(.2 <= p["lactate"] <= 30 for p in c["points"])
    assert all(a["power"] < b["power"] for a, b in zip(c["points"], c["points"][1:]))
    assert len(c["rest"]) == 1
(HERE / "inputs.json").write_text(json.dumps(curves, ensure_ascii=False, indent=2) + "\n")
print(f"Extracted {len(curves)} curves with source row references.")
